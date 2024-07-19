from ast import For
import json
from urllib import response
import pandas as pd
from openpyxl import load_workbook
import requests
from requests.auth import HTTPBasicAuth


url="https://maersk.service-now.com/api/now/v1/table/incident"
parameters= {
            "sysparm_display_value": True,
            "sysparm_exclude_reference_link": True,
            "sysparm_limit":50,
            "sysparm_query": "assignment_group=7a0a1c244fc1fa405f1db7718110c75b^ORassignment_group=b03a3c3f471539d0dd02e929736d43f1^state=1^ORstate=2^sys_updated_onONToday@javascript:gs.beginningOfToday()@javascript:gs.endOfToday()"
        }
response=requests.get(url=url,params=parameters,auth = HTTPBasicAuth('Tech_Support', ':hT}@_whX@s8Jhn$j6NXtr)^Q!{ePJ,;QvdV]1S6'))






# obj = json.loads(json_formatted_str)
# print(obj['result']['assignment_group'])
# print(obj['result'][0])

# print(obj['result'][1]['assigned_to'])

# List Records:
if response.status_code == 200:
    response_data=response.json()
    if response_data['result']:
        # json_formatted_str = json.dumps(response_data , indent=4)
        json_formatted_str = json.dumps(response_data)

    else:
        print("No Record found")
else:
    print("Error:", response.json())


# Get Excel data for LNS table


file_path = r'C:\Users\MSE230\Maersk Group\Test BOT group - Tushar-MBM\MBM-MTTD.xlsx'
sheet_name = 'LNS'
table_name = 'Table1'

df_sheet = pd.read_excel(file_path, sheet_name=sheet_name, engine='openpyxl')

# Load the workbook and specify the sheet
workbook = load_workbook(filename=file_path, data_only=True)
sheet = workbook[sheet_name]

# Access the table
table = sheet.tables[table_name]

# Get the range of the table
table_range = table.ref
start_cell, end_cell = table_range.split(':')
data = sheet[start_cell:end_cell]

# Extract the headers
headers = [cell.value for cell in data[0]]

# Extract the data
table_data = []
for row in data[1:]:
    table_data.append([cell.value for cell in row])

# Create a DataFrame for the table
df_table = pd.DataFrame(table_data, columns=headers)
 


# Loop through tickets fetched

for ticket in response_data['result']:
    # print(ticket['number'],':',ticket['state'])
    if df_table['ID'].isin([ticket['number']]).any():
        continue
    else:
        df_table.add



